"""批量导入学习资料目录中的可提取文本文件。"""
from pathlib import Path
import sys
import tempfile
import zipfile
import shutil

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")
import db
from rag import RAG

ROOT = Path(r"D:\新建文件夹 (5)\动物生理学复习题及答案(附解析)等99项文件")
SUPPORTED = {".docx", ".pptx", ".xlsx", ".xlsm", ".txt", ".md", ".csv", ".pdf"}

def extract(path):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv"}: return path.read_text(encoding="utf-8-sig", errors="ignore")
    if suffix == ".docx":
        from docx import Document
        doc = Document(path); rows = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables: rows.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(rows)
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".pptx":
        from pptx import Presentation
        rows = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip(): rows.append(shape.text.strip())
        return "\n".join(rows)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True); rows = []
        for ws in wb.worksheets:
            rows.append(f"【工作表：{ws.title}】")
            for row in ws.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values: rows.append(" | ".join(values))
        return "\n".join(rows)
    return ""

def main():
    if not ROOT.exists(): raise SystemExit(f"目录不存在：{ROOT}")
    rag = RAG(); conn = db.get_conn()
    existing = {row[0] for row in conn.execute("SELECT source FROM knowledge WHERE ktype='学习资料'")}
    imported = skipped = failed = 0
    # ZIP 中的资料先解压到临时目录，再与普通文件统一处理。
    unpack_root = Path(tempfile.mkdtemp(prefix="campus_zip_"))
    try:
        for archive in sorted(ROOT.rglob("*.zip")):
            target = unpack_root / archive.relative_to(ROOT).with_suffix("")
            try:
                with zipfile.ZipFile(archive) as z:
                    for member in z.infolist():
                        # 防止 ZIP 路径穿越
                        out = (target / member.filename).resolve()
                        if not str(out).startswith(str(target.resolve())):
                            continue
                        if not member.is_dir():
                            out.parent.mkdir(parents=True, exist_ok=True)
                            with z.open(member) as src, out.open("wb") as dst:
                                dst.write(src.read())
            except (zipfile.BadZipFile, OSError) as error:
                print(f"压缩包跳过：{archive} -> {error}", flush=True)
        scan_roots = [ROOT, unpack_root]
        all_paths = sorted(p for base in scan_roots for p in base.rglob("*") if p.is_file())
        for path in all_paths:
            if not path.is_file() or path.suffix.lower() not in SUPPORTED: continue
            if path.stat().st_size > 5 * 1024 * 1024:
                skipped += 1; continue
            if path.suffix.lower() == ".pdf" and path.stat().st_size > 10 * 1024 * 1024:
                skipped += 1; continue
            try:
                rel = path.relative_to(ROOT)
            except ValueError:
                rel = Path("压缩包解包") / path.relative_to(unpack_root)
            source = f"学习区文件：{rel}"
            if source in existing:
                skipped += 1; continue
            try:
                text = extract(path).strip()
                if len(text) < 30: skipped += 1; continue
                rag.add_document(path.stem[:180], text[:120000], source=source, ktype="学习资料", school_id="hlby", owner_user_id=None)
                existing.add(source)
                imported += 1
                if imported % 50 == 0: print(f"已导入 {imported} 个文件", flush=True)
            except Exception as error:
                failed += 1; print(f"跳过：{path} -> {error}", flush=True)
        print({"imported": imported, "skipped": skipped, "failed": failed})
    finally:
        shutil.rmtree(unpack_root, ignore_errors=True)

if __name__ == "__main__": main()
